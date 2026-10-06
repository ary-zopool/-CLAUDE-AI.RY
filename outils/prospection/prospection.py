"""
Robot de prospection LinkedIn — Ary (Suraya)

Utilisation (dans le dossier du robot) :
  python prospection.py --verifier     contrôle config, modèles et fichier Excel, sans navigateur
  python prospection.py --connexion    ouvre le navigateur pour vous connecter à LinkedIn (une seule fois)
  python prospection.py --essai        fait tout le parcours SANS rien envoyer (à faire avant le premier vrai lancement)
  python prospection.py                lancement normal (c'est ce que lance le Planificateur de tâches)
  Option : --forcer                    ignore les horaires et le retard au départ (pour vos tests)

Ce robot ne lit jamais les recherches LinkedIn : il ne traite que les profils que vous avez ajoutés dans prospects.xlsx.
Il s'arrête net si LinkedIn affiche une vérification ou un avertissement.
"""

import argparse
import datetime as dt
import logging
import random
import re
import smtplib
import sys
import time
from email.message import EmailMessage
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

DOSSIER = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# Sélecteurs LinkedIn : SEUL endroit à corriger quand LinkedIn change son interface.
# ---------------------------------------------------------------------------
SELECTEURS = {
    "bouton_connecter": re.compile(r"(Invitez .* à rejoindre votre réseau|Invite .* to connect|^Se connecter$|^Connect$)", re.I),
    "bouton_plus": re.compile(r"^(Plus|More|Plus d.actions|More actions)$", re.I),
    "menu_connecter": re.compile(r"(Se connecter|Connect|Invitez .* à rejoindre)", re.I),
    "envoyer_sans_note": re.compile(r"(Envoyer sans note|Send without a note)", re.I),
    "bouton_en_attente": re.compile(r"(En attente|Pending)", re.I),
    "bouton_message": re.compile(r"^(Message|Envoyer un message)", re.I),
    "zone_saisie": "div.msg-form__contenteditable[contenteditable='true']",
    "bouton_envoyer": "button.msg-form__send-button",
    "message_de_l_autre": ".msg-s-event-listitem--other",
    "fermer_conversation": "button.msg-overlay-bubble-header__control:has(svg[data-test-icon='close-small'])",
}
SIGNES_DE_BLOCAGE_URL = ("/checkpoint/", "/authwall", "/login", "/uas/")
SIGNES_DE_BLOCAGE_TEXTE = re.compile(
    r"(activité inhabituelle|unusual activity|temporairement restreint|temporarily restricted|"
    r"vérifi(ons|er) que vous (n.êtes pas un robot|êtes humain)|security verification|limite hebdomadaire|weekly invitation limit)",
    re.I,
)

# Mots interdits côté EZdrive (charte V3, skill ligne-editoriale)
MOTS_INTERDITS_EZDRIVE = [
    "recharge solaire", "rechargement solaire", "recharge à l'énergie solaire", "énergie solaire",
    "borne solaire", "bornes solaires", "borne photovoltaïque", "bornes photovoltaïques",
    "production solaire", "premier réseau solaire", "ezsolar", "intelligence artificielle",
    "station de charge", "métropole", "leader", "n°1", "pionnier", "de référence",
]
MOTS_INTERDITS_EXACTS = ["IA", "AI"]          # mots entiers seulement
TUTOIEMENT = re.compile(r"\b(tu|toi|ta|tes|ton)\b", re.I)

COLONNES = [
    "url_profil", "prenom", "entreprise", "poste", "territoire", "secteur",
    "statut", "date_invitation", "date_message_1", "date_relance_1", "date_relance_2",
    "date_reponse", "derniere_erreur", "notes",
]
STATUTS = {
    "a_inviter": "à ajouter par vous",
    "invite": "invitation envoyée, en attente d'acceptation",
    "message_1": "message 1 envoyé",
    "relance_1": "relance 1 envoyée",
    "relance_2": "relance 2 envoyée",
    "termine": "séquence finie sans réponse",
    "repondu": "a répondu : à vous de jouer",
    "exclu": "ne pas contacter",
    "erreur": "problème : voir derniere_erreur",
}
JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]


class BlocageLinkedIn(Exception):
    """LinkedIn affiche une vérification ou un avertissement : on arrête tout."""


# ---------------------------------------------------------------------------
# Chargement
# ---------------------------------------------------------------------------
def charger_yaml(nom):
    with open(DOSSIER / nom, encoding="utf-8") as f:
        return yaml.safe_load(f)


def configurer_journal(config):
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)s  %(message)s",
        handlers=[
            logging.FileHandler(DOSSIER / config["fichier_journal"], encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )


def verifier_charte(texte, nom):
    """Renvoie la liste des problèmes trouvés dans un modèle."""
    problemes = []
    bas = texte.lower()
    for mot in MOTS_INTERDITS_EZDRIVE:
        if mot in bas:
            problemes.append(f"{nom} : mot interdit « {mot} »")
    for mot in MOTS_INTERDITS_EXACTS:
        if re.search(rf"\b{mot}\b", texte):
            problemes.append(f"{nom} : mot interdit « {mot} »")
    if TUTOIEMENT.search(texte):
        problemes.append(f"{nom} : tutoiement détecté (vouvoiement obligatoire)")
    if "http" in bas or "www." in bas:
        problemes.append(f"{nom} : lien dans le message (à éviter, LinkedIn freine ces messages)")
    return problemes


def charger_campagne(config):
    modeles = charger_yaml("modeles.yaml")
    nom = config["campagne_active"]
    if nom not in (modeles or {}):
        raise SystemExit(f"Campagne « {nom} » absente de modeles.yaml.")
    campagne = modeles[nom]
    for cle in ("message_1", "relance_1", "relance_2"):
        if not campagne.get(cle):
            raise SystemExit(f"Campagne « {nom} » : le modèle {cle} est vide.")
    if campagne.get("verifier_charte"):
        problemes = []
        for cle in ("message_1", "relance_1", "relance_2"):
            problemes += verifier_charte(campagne[cle], cle)
        if problemes:
            raise SystemExit("Modèles refusés (charte V3) :\n  - " + "\n  - ".join(problemes))
    return campagne


def remplir(modele, ligne):
    """Remplit un modèle. Lève ValueError si une variable utilisée est vide."""
    variables = set(re.findall(r"{(\w+)}", modele))
    valeurs = {}
    for v in variables:
        val = str(ligne.get(v) or "").strip()
        if not val:
            raise ValueError(f"variable vide : {v}")
        valeurs[v] = val
    return modele.format(**valeurs)


# ---------------------------------------------------------------------------
# Fichier Excel
# ---------------------------------------------------------------------------
def creer_fichier_prospects(chemin):
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    gras = Font(name="Arial", bold=True)
    normal = Font(name="Arial")
    saisie = PatternFill("solid", start_color="FFF2CC")
    ws.append(COLONNES)
    for c in ws[1]:
        c.font = gras
    exemple = ["https://www.linkedin.com/in/exemple-profil/", "Marie", "Hôtel Exemple", "Directrice",
               "Martinique", "Hôtellerie", "exclu", "", "", "", "", "", "", "LIGNE D'EXEMPLE : à supprimer"]
    ws.append(exemple)
    for c in ws[2]:
        c.font = normal
    for col in range(1, 7):
        ws.cell(row=1, column=col).fill = saisie
    for i, largeur in enumerate([45, 14, 24, 22, 14, 16, 14, 16, 16, 16, 16, 16, 30, 30], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = largeur
    ws.freeze_panes = "A2"

    leg = wb.create_sheet("Mode d'emploi")
    lignes = [
        ["Mode d'emploi"],
        [""],
        ["Vous remplissez", "Les colonnes en jaune : url_profil, prenom, entreprise, poste, territoire, secteur."],
        ["Statut à mettre", "a_inviter pour lancer un prospect. exclu pour ne jamais le contacter."],
        ["Le robot remplit", "statut et toutes les dates. Ne les modifiez pas, sauf pour remettre un prospect en a_inviter."],
        ["Fichier ouvert", "Fermez ce fichier avant que le robot tourne, sinon il ne peut pas enregistrer."],
        [""],
        ["Statut", "Signification"],
    ] + [[k, v] for k, v in STATUTS.items()]
    for r in lignes:
        leg.append(r)
    for row in leg.iter_rows():
        for c in row:
            c.font = normal
    leg["A1"].font = Font(name="Arial", bold=True, size=14)
    leg["A8"].font = gras
    leg["B8"].font = gras
    leg.column_dimensions["A"].width = 18
    leg.column_dimensions["B"].width = 95
    wb.save(chemin)


class Prospects:
    def __init__(self, chemin):
        self.chemin = chemin
        if not chemin.exists():
            creer_fichier_prospects(chemin)
            logging.info("Fichier %s créé. Remplissez-le, puis relancez.", chemin.name)
        self.wb = load_workbook(chemin)
        self.ws = self.wb["Prospects"]
        entetes = [c.value for c in self.ws[1]]
        manquantes = [c for c in COLONNES if c not in entetes]
        if manquantes:
            raise SystemExit(f"Colonnes manquantes dans {chemin.name} : {', '.join(manquantes)}")
        self.index = {nom: entetes.index(nom) + 1 for nom in COLONNES}

    def lignes(self):
        for r in range(2, self.ws.max_row + 1):
            ligne = {nom: self.ws.cell(row=r, column=i).value for nom, i in self.index.items()}
            if ligne["url_profil"]:
                ligne["_ligne"] = r
                ligne["statut"] = str(ligne["statut"] or "").strip()
                yield ligne

    def maj(self, ligne, **valeurs):
        for cle, val in valeurs.items():
            self.ws.cell(row=ligne["_ligne"], column=self.index[cle], value=val)
            ligne[cle] = val
        self.enregistrer()

    def enregistrer(self):
        try:
            self.wb.save(self.chemin)
        except PermissionError:
            raise SystemExit(f"Impossible d'enregistrer {self.chemin.name} : fermez-le dans Excel puis relancez.")


def aujourd_hui():
    return dt.date.today().isoformat()


def jours_depuis(date_texte):
    if not date_texte:
        return 0
    if isinstance(date_texte, dt.datetime):
        d = date_texte.date()
    elif isinstance(date_texte, dt.date):
        d = date_texte
    else:
        d = dt.date.fromisoformat(str(date_texte)[:10])
    return (dt.date.today() - d).days


def est_exclu(ligne, config):
    entreprise = str(ligne.get("entreprise") or "").lower()
    for e in config.get("exclusions_entreprises") or []:
        if e and e.lower() in entreprise:
            return True
    url = str(ligne.get("url_profil") or "").rstrip("/").lower()
    return url in [str(u).rstrip("/").lower() for u in (config.get("exclusions_profils") or [])]


# ---------------------------------------------------------------------------
# Alertes
# ---------------------------------------------------------------------------
def alerter(config, sujet, corps):
    logging.warning("ALERTE : %s — %s", sujet, corps)
    a = config.get("alerte_email") or {}
    if not a.get("actif"):
        return
    try:
        msg = EmailMessage()
        msg["Subject"] = f"[Prospection] {sujet}"
        msg["From"] = a["expediteur"]
        msg["To"] = a["destinataire"]
        msg.set_content(corps)
        with smtplib.SMTP(a["smtp_serveur"], int(a["smtp_port"])) as s:
            s.starttls()
            s.login(a["expediteur"], a["mot_de_passe_application"])
            s.send_message(msg)
    except Exception as e:  # une alerte ratée ne doit pas arrêter le robot
        logging.error("Alerte e-mail non envoyée : %s", e)


# ---------------------------------------------------------------------------
# LinkedIn
# ---------------------------------------------------------------------------
class LinkedIn:
    def __init__(self, page, config, essai):
        self.page = page
        self.config = config
        self.essai = essai
        self.visites = 0

    def pause(self, facteur=1.0):
        time.sleep(random.uniform(self.config["pause_min_secondes"], self.config["pause_max_secondes"]) * facteur)

    def petite_pause(self):
        time.sleep(random.uniform(1.5, 4.0))

    def controler_blocage(self):
        url = self.page.url
        if any(s in url for s in SIGNES_DE_BLOCAGE_URL):
            raise BlocageLinkedIn(f"page de vérification ou de connexion : {url}")
        try:
            texte = self.page.inner_text("body", timeout=5000)
        except Exception:
            return
        m = SIGNES_DE_BLOCAGE_TEXTE.search(texte)
        if m:
            raise BlocageLinkedIn(f"message LinkedIn détecté : « {m.group(0)} »")

    def quota_visites_atteint(self):
        return self.visites >= self.config["visites_profil_par_jour"]

    def ouvrir_profil(self, url):
        self.visites += 1
        self.page.goto(url, wait_until="domcontentloaded", timeout=45000)
        self.page.wait_for_timeout(random.randint(3000, 6000))
        self.controler_blocage()

    def bouton(self, motif):
        """Premier bouton visible de la zone principale du profil dont le nom correspond."""
        boutons = self.page.locator("main").get_by_role("button", name=motif)
        for i in range(min(boutons.count(), 6)):
            b = boutons.nth(i)
            if b.is_visible():
                return b
        return None

    def etat_relation(self):
        """Renvoie 'connecte', 'en_attente' ou 'non_connecte'."""
        if self.bouton(SELECTEURS["bouton_en_attente"]):
            return "en_attente"
        if self.bouton(SELECTEURS["bouton_connecter"]):
            return "non_connecte"
        if self.bouton(SELECTEURS["bouton_message"]):
            return "connecte"
        return "non_connecte"

    def inviter(self):
        b = self.bouton(SELECTEURS["bouton_connecter"])
        if not b:
            plus = self.bouton(SELECTEURS["bouton_plus"])
            if not plus:
                raise RuntimeError("bouton « Se connecter » introuvable")
            plus.click()
            self.petite_pause()
            items = self.page.get_by_role("menuitem", name=SELECTEURS["menu_connecter"])
            if items.count() == 0:
                items = self.page.locator("div[role='button']", has_text=SELECTEURS["menu_connecter"])
            if items.count() == 0:
                self.page.keyboard.press("Escape")
                raise RuntimeError("« Se connecter » absent du menu Plus")
            b = items.first
        if self.essai:
            logging.info("   [essai] invitation non envoyée")
            self.page.keyboard.press("Escape")
            return
        b.click()
        self.petite_pause()
        sans_note = self.page.get_by_role("button", name=SELECTEURS["envoyer_sans_note"])
        if sans_note.count() == 0:
            self.page.keyboard.press("Escape")
            raise RuntimeError("bouton « Envoyer sans note » introuvable")
        sans_note.first.click()
        self.page.wait_for_timeout(2500)
        self.controler_blocage()

    def ouvrir_conversation(self):
        b = self.bouton(SELECTEURS["bouton_message"])
        if not b:
            raise RuntimeError("bouton « Message » introuvable")
        b.click()
        self.page.wait_for_selector(SELECTEURS["zone_saisie"], timeout=15000)
        self.page.wait_for_timeout(2500)
        self.controler_blocage()

    def a_repondu(self):
        return self.page.locator(SELECTEURS["message_de_l_autre"]).count() > 0

    def ecrire_et_envoyer(self, texte):
        zone = self.page.locator(SELECTEURS["zone_saisie"]).last
        zone.click()
        lignes = texte.split("\n")
        for i, l in enumerate(lignes):
            zone.type(l, delay=random.randint(25, 70))
            if i < len(lignes) - 1:
                self.page.keyboard.press("Shift+Enter")
        self.petite_pause()
        if self.essai:
            logging.info("   [essai] message non envoyé")
            zone.fill("")
        else:
            envoyer = self.page.locator(SELECTEURS["bouton_envoyer"]).last
            if not envoyer.is_enabled():
                raise RuntimeError("bouton « Envoyer » inactif")
            envoyer.click()
            self.page.wait_for_timeout(2500)
        self.fermer_conversation()

    def fermer_conversation(self):
        try:
            for b in self.page.locator(SELECTEURS["fermer_conversation"]).all():
                b.click(timeout=2000)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Déroulé d'une journée
# ---------------------------------------------------------------------------
def dans_les_horaires(config):
    maintenant = dt.datetime.now()
    jour = JOURS[maintenant.weekday()]
    return jour in config["jours_autorises"] and config["heure_debut"] <= maintenant.hour < config["heure_fin"]


def traiter_suivis(li, prospects, campagne, config, compteurs):
    """Acceptations, réponses et relances pour les prospects déjà invités."""
    for ligne in list(prospects.lignes()):
        if compteurs["messages"] >= config["messages_par_jour"] or li.quota_visites_atteint():
            break
        if not dans_les_horaires(config) and not compteurs["forcer"]:
            break
        statut = ligne["statut"]
        if statut not in ("invite", "message_1", "relance_1", "relance_2"):
            continue
        if statut == "invite" and jours_depuis(ligne["date_invitation"]) > 21:
            prospects.maj(ligne, statut="termine", derniere_erreur="invitation non acceptée après 21 jours")
            continue
        if statut == "relance_2":
            if jours_depuis(ligne["date_relance_2"]) >= 7:
                prospects.maj(ligne, statut="termine")
            continue
        if statut == "message_1" and jours_depuis(ligne["date_message_1"]) < config["relance_1_jours"]:
            continue
        if statut == "relance_1" and jours_depuis(ligne["date_message_1"]) < config["relance_2_jours"]:
            continue

        nom = f"{ligne['prenom']} ({ligne['entreprise']})"
        try:
            li.ouvrir_profil(ligne["url_profil"])
            etat = li.etat_relation()
            if statut == "invite":
                if etat != "connecte":
                    continue
                texte = remplir(campagne["message_1"], ligne)
                li.ouvrir_conversation()
                if li.a_repondu():
                    li.fermer_conversation()
                    prospects.maj(ligne, statut="repondu", date_reponse=aujourd_hui())
                    alerter(config, f"Réponse de {nom}", f"{nom} vous a écrit sur LinkedIn.\n{ligne['url_profil']}")
                    continue
                logging.info("Message 1 → %s", nom)
                li.ecrire_et_envoyer(texte)
                if not li.essai:
                    prospects.maj(ligne, statut="message_1", date_message_1=aujourd_hui(), derniere_erreur="")
            else:
                prochain = "relance_1" if statut == "message_1" else "relance_2"
                texte = remplir(campagne[prochain], ligne)
                li.ouvrir_conversation()
                if li.a_repondu():
                    li.fermer_conversation()
                    prospects.maj(ligne, statut="repondu", date_reponse=aujourd_hui())
                    alerter(config, f"Réponse de {nom}", f"{nom} vous a répondu sur LinkedIn.\n{ligne['url_profil']}")
                    continue
                logging.info("%s → %s", prochain.replace("_", " ").capitalize(), nom)
                li.ecrire_et_envoyer(texte)
                if not li.essai:
                    prospects.maj(ligne, statut=prochain, **{f"date_{prochain}": aujourd_hui()}, derniere_erreur="")
            compteurs["messages"] += 1
            li.pause()
        except BlocageLinkedIn:
            raise
        except ValueError as e:
            prospects.maj(ligne, statut="erreur", derniere_erreur=str(e))
            logging.error("%s : %s", nom, e)
        except Exception as e:
            prospects.maj(ligne, derniere_erreur=f"{aujourd_hui()} : {str(e)[:150]}")
            logging.error("%s : %s", nom, e)
            li.fermer_conversation()
            li.pause(0.5)


def traiter_invitations(li, prospects, campagne, config, compteurs):
    for ligne in list(prospects.lignes()):
        if compteurs["invitations"] >= config["invitations_par_jour"] or li.quota_visites_atteint():
            break
        if not dans_les_horaires(config) and not compteurs["forcer"]:
            break
        if ligne["statut"] != "a_inviter":
            continue
        nom = f"{ligne['prenom']} ({ligne['entreprise']})"
        if est_exclu(ligne, config):
            prospects.maj(ligne, statut="exclu")
            continue
        try:
            remplir(campagne["message_1"], ligne)   # on vérifie dès maintenant que le message 1 pourra partir
        except ValueError as e:
            prospects.maj(ligne, statut="erreur", derniere_erreur=str(e))
            logging.error("%s : %s", nom, e)
            continue
        try:
            li.ouvrir_profil(ligne["url_profil"])
            etat = li.etat_relation()
            if etat == "connecte":
                logging.info("Déjà en relation → %s : passe directement au message 1", nom)
                if not li.essai:
                    prospects.maj(ligne, statut="invite", date_invitation=aujourd_hui())
                continue
            if etat == "en_attente":
                if not li.essai:
                    prospects.maj(ligne, statut="invite", date_invitation=aujourd_hui())
                continue
            logging.info("Invitation → %s", nom)
            li.inviter()
            if not li.essai:
                prospects.maj(ligne, statut="invite", date_invitation=aujourd_hui(), derniere_erreur="")
            compteurs["invitations"] += 1
            li.pause()
        except BlocageLinkedIn:
            raise
        except Exception as e:
            prospects.maj(ligne, derniere_erreur=f"{aujourd_hui()} : {str(e)[:150]}")
            logging.error("%s : %s", nom, e)
            li.pause(0.5)


def lancer(config, campagne, essai, forcer, connexion):
    from playwright.sync_api import sync_playwright   # importé ici pour que --verifier marche sans navigateur

    with sync_playwright() as p:
        contexte = p.chromium.launch_persistent_context(
            str(DOSSIER / config["dossier_navigateur"]),
            headless=False,
            locale="fr-FR",
            viewport={"width": 1366, "height": 850},
        )
        page = contexte.pages[0] if contexte.pages else contexte.new_page()
        if connexion:
            page.goto("https://www.linkedin.com/login")
            print("\nConnectez-vous à LinkedIn dans la fenêtre ouverte, puis revenez ici et appuyez sur Entrée.")
            input()
            contexte.close()
            print("Connexion enregistrée. Lancez maintenant : python prospection.py --essai --forcer")
            return

        prospects = Prospects(DOSSIER / config["fichier_prospects"])
        li = LinkedIn(page, config, essai)
        compteurs = {"invitations": 0, "messages": 0, "forcer": forcer}
        try:
            page.goto("https://www.linkedin.com/feed/", wait_until="domcontentloaded", timeout=45000)
            page.wait_for_timeout(4000)
            li.controler_blocage()
            traiter_suivis(li, prospects, campagne, config, compteurs)
            traiter_invitations(li, prospects, campagne, config, compteurs)
            logging.info("Fin : %s invitation(s), %s message(s)%s.",
                         compteurs["invitations"], compteurs["messages"], " (essai, rien envoyé)" if essai else "")
        except BlocageLinkedIn as e:
            alerter(config, "ROBOT ARRÊTÉ — LinkedIn demande une vérification",
                    f"{e}\n\nNe relancez pas le robot. Ouvrez LinkedIn vous-même, répondez à la vérification, "
                    f"et attendez au moins 48 h avant de relancer.")
            sys.exit(2)
        finally:
            prospects.enregistrer()
            contexte.close()


def main():
    parser = argparse.ArgumentParser(description="Robot de prospection LinkedIn")
    parser.add_argument("--verifier", action="store_true")
    parser.add_argument("--connexion", action="store_true")
    parser.add_argument("--essai", action="store_true")
    parser.add_argument("--forcer", action="store_true")
    args = parser.parse_args()

    config = charger_yaml("config.yaml")
    configurer_journal(config)
    if config["invitations_par_jour"] > 20:
        raise SystemExit("invitations_par_jour au-dessus de 20 : trop risqué pour votre profil.")
    campagne = charger_campagne(config)

    if args.verifier:
        prospects = Prospects(DOSSIER / config["fichier_prospects"])
        n = {}
        for l in prospects.lignes():
            n[l["statut"] or "(vide)"] = n.get(l["statut"] or "(vide)", 0) + 1
        logging.info("Vérification OK. Campagne : %s. Prospects par statut : %s", config["campagne_active"], n or "aucun")
        return

    if args.connexion:
        lancer(config, campagne, essai=True, forcer=True, connexion=True)
        return

    if not args.forcer:
        if not dans_les_horaires(config):
            logging.info("Hors horaires autorisés : rien à faire.")
            return
        attente = random.randint(0, config["retard_depart_max_minutes"] * 60)
        logging.info("Départ dans %s min.", attente // 60)
        time.sleep(attente)

    lancer(config, campagne, essai=args.essai, forcer=args.forcer, connexion=False)


if __name__ == "__main__":
    main()
